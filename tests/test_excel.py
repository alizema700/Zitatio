"""Excel-Mappe mit Formeln: lesbar, und die Rechnungen stehen als Formeln drin."""

import io
from pathlib import Path

import openpyxl

from zitatio.dataroom import DataRoom
from zitatio.excel import to_xlsx
from zitatio.pipeline import run_check
from zitatio.report import load_report

GROSS = Path(__file__).resolve().parents[1] / "beispiele" / "gross"


def test_workbook_has_sheets_and_formulas():
    text, data = load_report(GROSS / "dd-bericht.md")
    run = run_check("dd-bericht.md", text, data, DataRoom.from_folder(GROSS / "datenraum"))
    wb = openpyxl.load_workbook(io.BytesIO(to_xlsx(run)))
    assert wb.sheetnames == ["Aussagen", "EBITDA-Brücke", "Nettoverschuldung", "Closing-Liste", "Datenraum"]
    claims = wb["Aussagen"]
    assert claims.max_row == len(run.findings) + 1 and claims["E1"].value == "Status"
    formulas = [c.value for ws in wb for row in ws.iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("=")]
    assert any(f.startswith("=SUM(") for f in formulas)
    assert any("SUMIF" in f for f in formulas) and any(f.startswith('=IF(D') and '"ja"' in f for f in formulas)
