import io
import zipfile
from decimal import Decimal

from openpyxl import Workbook

from zitatio.dataroom import DataRoom
from zitatio.facts import AMOUNT, PERCENT, Fact, extract_facts, page_facts, page_unit
from zitatio.pipeline import run_check
from zitatio.render import to_html
from zitatio.report import parse_report
from zitatio.verify import Status, verify_claim
from zitatio.xlsx import xlsx_to_pages


def amount(value):
    return Fact(AMOUNT, ("CHF", Decimal(value)), "", 0, 0)


def test_rounded_amount_matches_within_precision_and_one_percent():
    claim = extract_facts("CHF 18.45 Mio.")[0]
    assert claim.match(amount(18_450_312)) == "gerundet"
    assert claim.match(amount(18_470_000)) is None
    assert extract_facts("CHF 3 Mio.")[0].match(amount(3_400_000)) is None  # 13 % daneben


def test_whole_percent_is_not_rounded():
    one_decimal = extract_facts("12.4 %")[0]
    assert one_decimal.match(Fact(PERCENT, Decimal("12.38"), "", 0, 0)) == "gerundet"
    assert extract_facts("2 %")[0].match(Fact(PERCENT, Decimal("2.4"), "", 0, 0)) is None


def test_page_unit_needs_one_unambiguous_header():
    assert page_unit("Erfolgsrechnung (in TCHF)") == ("CHF", 1000)
    assert page_unit("Bilanz in Mio. CHF") == ("CHF", 10**6)
    assert page_unit("Kaufpreis CHF 3 Mio. in bar") is None
    assert page_unit("in CHF\nVorjahr in TCHF") is None


def test_bare_table_numbers_take_the_header_unit():
    values = [f.value for f in page_facts("Erfolgsrechnung 2025 (in TCHF)\nNettoerlös 18'450\nAnmerkung 3")]
    assert ("CHF", Decimal(18_450_000)) in values
    assert all(v != ("CHF", Decimal(2_025_000)) and v != ("CHF", Decimal(3000)) for v in values)


def _workbook() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "ER"
    ws.append(["Erfolgsrechnung in TCHF", 2025])
    ws.append(["Nettoerlös", 18450.3])
    ws["B2"].number_format = "#,##0"
    ws.append(["Marge", 0.1214])
    ws["B3"].number_format = "0.0%"
    wb.create_sheet("Zweites Blatt").append(["Kasse", 1200])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_xlsx_sheets_become_pages_with_displayed_values():
    pages = xlsx_to_pages(_workbook())
    assert len(pages) == 2
    assert "Nettoerlös; 18'450" in pages[0] and "Marge; 12.1 %" in pages[0]
    assert pages[1].startswith("Blatt: Zweites Blatt")


def test_zip_data_room_reads_subfolders_and_normalizes_index():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("Deal/02 Finanzen/02.2 Plan.xlsx", _workbook())
        z.writestr("Deal/.DS_Store", b"x")
        z.writestr("__MACOSX/Deal/._02.2 Plan.xlsx", b"x")
    room = DataRoom()
    room.add("datenraum.zip", buf.getvalue())
    assert [d.index for d in room.documents] == ["2.2"]
    claim = parse_report("Der Nettoerlös beträgt CHF 18.45 Mio. [DR 2.2, S. 1].")[0]
    finding = verify_claim(claim, room)
    assert finding.status is Status.GEPRUEFT
    assert "Tabelle: 18'450 TCHF" in finding.reasons[0]


def test_report_has_review_controls():
    room = DataRoom()
    room.add("1 Plan.txt", b"Umsatz CHF 1'000'000")
    text = "Umsatz CHF 1 Mio. [DR 1]."
    page = to_html(run_check("b.md", text, text.encode(), room))
    assert 'class="decide"' in page and 'id="reviewer"' in page and "Protokoll exportieren" in page


def test_adjacent_table_columns_are_separate_numbers():
    values = [f.value for f in page_facts("Erfolgsrechnung (in CHF)\nFinanzaufwand 187'400 201'300")]
    assert values == [("CHF", Decimal(187_400)), ("CHF", Decimal(201_300))]


def test_sentence_starting_with_year_after_citation():
    claims = parse_report("Umsatz CHF 2 Mio. [DR 1, S. 1]. 2023 waren es CHF 1 Mio. [DR 1, S. 2].")
    assert len(claims) == 2
